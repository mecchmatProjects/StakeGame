from __future__ import annotations

import csv
import json
import math
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
import zipfile
from typing import Any, Dict, Iterable, List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(r"f:/Gaming/Stakes/math-sdk/games/treasure_dice_final")
PUBLISH_DIR = ROOT / "artifacts" / "publish_files"
CONFIG_DIR = ROOT / "artifacts" / "configs"
MAPPING_DIR = ROOT / "artifacts" / "mapping"
DOCS_DIR = ROOT / "docs"
OUTPUT = ROOT / "TreasureDiceFull.xlsx"
PAYOUT_SCALE = 100.0
SMALL_BETS = [0.01, 0.02, 0.05]

DOC_FILES = [
    "Treasure_Dice_Final_Math_Report.md",
    "Treasure_Dice_Approved_Regular_Routes_Paytable.md",
    "Treasure_Dice_Approved_Bonus_Buy_Paytable.md",
    "Treasure_Dice_Phase_3_Event_Mapping_Report.md",
    "Treasure_Dice_Phase_5_Simulation_Report.md",
    "Treasure_Dice_Phase_5_Volatility_Report.md",
    "Treasure_Dice_Phase_5_Final_Validation_Summary.md",
    "Treasure_Dice_Phase_6_Precision_Rounding_And_Frontend_Contract.md",
    "Treasure_Dice_Exposure_Table_Pending_Operator_Cap.md",
    "Treasure_Dice_Phase_6_Risk_Assessment.md",
    "Treasure_Dice_Phase_6_Acceptance_Checklist.md",
    "Treasure_Dice_Phase_6_Publication_Status_Note.md",
]

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FILL = PatternFill("solid", fgColor="B4C6E7")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")

TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)
HEADER_FONT = Font(bold=True)
BOLD_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


@dataclass
class LookupStats:
    mode: str
    cost: float
    total_weight: int
    weighted_mean: float
    rtp: float
    hit_rate: float
    zero_rate: float
    below_cost_rate: float
    profit_rate: float
    variance: float
    std: float
    min_win: float
    max_win: float
    median: float
    p95: float
    p99: float
    rows: List[Tuple[int, int, float, float, float, int, float, float]]


@dataclass
class UniqueOutcomeStats:
    payout: float
    weight: int
    probability: float
    weighted_contribution: float
    cumulative_probability: float


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def safe_sheet_name(name: str) -> str:
    cleaned = "".join(ch if ch.isalnum() else "_" for ch in name)
    return cleaned[:31]


def set_col_widths(ws, widths: Dict[int, float]) -> None:
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width



def add_title(ws, title: str, subtitle: str = "") -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = WRAP
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True)
        ws["A2"].alignment = WRAP


def section(ws, row: int, title: str) -> int:
    ws.cell(row, 1, title)
    ws.cell(row, 1).font = BOLD_FONT
    ws.cell(row, 1).fill = SECTION_FILL
    return row + 1


def kv_table(ws, start_row: int, items: List[Tuple[str, Any]], formats: Dict[str, str] | None = None) -> int:
    formats = formats or {}
    ws.cell(start_row, 1, "Key")
    ws.cell(start_row, 2, "Value")
    ws.cell(start_row, 1).fill = HEADER_FILL
    ws.cell(start_row, 2).fill = HEADER_FILL
    ws.cell(start_row, 1).font = HEADER_FONT
    ws.cell(start_row, 2).font = HEADER_FONT
    row = start_row + 1
    for key, value in items:
        ws.cell(row, 1, key)
        ws.cell(row, 2, value)
        if key in formats:
            ws.cell(row, 2).number_format = formats[key]
        row += 1
    return row


def data_table(ws, start_row: int, headers: List[str], rows: Iterable[Iterable[Any]]) -> int:
    for col, header in enumerate(headers, start=1):
        ws.cell(start_row, col, header)
        ws.cell(start_row, col).fill = HEADER_FILL
        ws.cell(start_row, col).font = HEADER_FONT
        ws.cell(start_row, col).alignment = WRAP
    row = start_row + 1
    for record in rows:
        for col, value in enumerate(record, start=1):
            ws.cell(row, col, value)
            ws.cell(row, col).alignment = TOP
        row += 1
    return row


def weighted_percentile(values: List[float], weights: List[int], pct: float) -> float:
    if not values:
        return 0.0
    total = sum(weights)
    if total <= 0:
        return 0.0
    ordered = sorted(zip(values, weights), key=lambda item: item[0])
    threshold = total * pct
    cumulative = 0
    for value, weight in ordered:
        cumulative += weight
        if cumulative >= threshold:
            return float(value)
    return float(ordered[-1][0])


def calc_lookup_stats(mode: str, cost: float, rows: List[Tuple[int, int, float]]) -> LookupStats:
    total_weight = sum(weight for _, weight, _ in rows)
    payouts = [raw / PAYOUT_SCALE for _, _, raw in rows]
    if total_weight <= 0:
        return LookupStats(mode, cost, 0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, [])

    probs = [weight / total_weight for _, weight, _ in rows]
    weighted_mean = sum(prob * payout for prob, payout in zip(probs, payouts))
    rtp = weighted_mean / cost
    hit_rate = sum(prob for prob, payout in zip(probs, payouts) if payout > 0)
    zero_rate = sum(prob for prob, payout in zip(probs, payouts) if payout == 0)
    below_cost_rate = sum(prob for prob, payout in zip(probs, payouts) if 0 < payout < cost)
    profit_rate = sum(prob for prob, payout in zip(probs, payouts) if payout > cost)
    variance = sum(prob * ((payout - weighted_mean) ** 2) for prob, payout in zip(probs, payouts))
    std = math.sqrt(variance)
    min_win = min(payouts)
    max_win = max(payouts)
    median = weighted_percentile(payouts, [weight for _, weight, _ in rows], 0.50)
    p95 = weighted_percentile(payouts, [weight for _, weight, _ in rows], 0.95)
    p99 = weighted_percentile(payouts, [weight for _, weight, _ in rows], 0.99)

    cumulative_weight = 0
    cumulative_prob = 0.0
    audit_rows: List[Tuple[int, int, float, float, float, int, float, float]] = []
    for sim_id, weight, raw in rows:
        payout = raw / PAYOUT_SCALE
        prob = weight / total_weight
        weighted_contribution = prob * payout
        cumulative_weight += weight
        cumulative_prob += prob
        audit_rows.append((sim_id, weight, raw, payout, prob, cumulative_weight, cumulative_prob, weighted_contribution))

    return LookupStats(
        mode=mode,
        cost=cost,
        total_weight=total_weight,
        weighted_mean=weighted_mean,
        rtp=rtp,
        hit_rate=hit_rate,
        zero_rate=zero_rate,
        below_cost_rate=below_cost_rate,
        profit_rate=profit_rate,
        variance=variance,
        std=std,
        min_win=min_win,
        max_win=max_win,
        median=median,
        p95=p95,
        p99=p99,
        rows=audit_rows,
    )


def unique_outcomes(rows: List[Tuple[int, int, float]]) -> List[UniqueOutcomeStats]:
    grouped: Dict[float, int] = defaultdict(int)
    for _, weight, raw in rows:
        grouped[raw / PAYOUT_SCALE] += weight
    total_weight = sum(grouped.values())
    cumulative = 0.0
    out: List[UniqueOutcomeStats] = []
    for payout in sorted(grouped):
        weight = grouped[payout]
        probability = weight / total_weight if total_weight else 0.0
        cumulative += probability
        out.append(UniqueOutcomeStats(payout, weight, probability, payout * probability, cumulative))
    return out


def load_lookup_rows(path: Path) -> List[Tuple[int, int, float]]:
    rows: List[Tuple[int, int, float]] = []
    with path.open("r", encoding="utf-8", newline="") as handle:
        for row in csv.reader(handle):
            if len(row) < 3:
                continue
            rows.append((int(float(row[0])), int(float(row[1])), float(row[2])))
    return rows


def write_lookup_sheet(ws, stats: LookupStats) -> None:
    add_title(ws, f"Lookup Audit - {stats.mode}", f"Source: lookUpTable_{stats.mode}_0.csv")
    ws.cell(4, 1, "Key")
    ws.cell(4, 2, "Value")
    ws.cell(4, 1).fill = HEADER_FILL
    ws.cell(4, 2).fill = HEADER_FILL
    ws.cell(4, 1).font = HEADER_FONT
    ws.cell(4, 2).font = HEADER_FONT

    metric_labels = [
        "Cost",
        "Total Weight",
        "Weighted Mean Payout",
        "RTP",
        "Hit Rate",
        "Zero Rate",
        "Below-Cost Rate",
        "Profit Rate",
        "Variance",
        "Std Dev",
        "Min Win",
        "Max Win",
        "Median",
        "P95",
        "P99",
    ]
    metric_rows = {label: 6 + idx for idx, label in enumerate(metric_labels)}
    for label, row_num in metric_rows.items():
        ws.cell(row_num, 1, label)
        ws.cell(row_num, 1).font = BOLD_FONT

    header_row = 22
    data_start = 23
    data_end = data_start + len(stats.rows) - 1
    headers = ["ID", "Weight", "Payout Raw", "Payout Norm", "Probability", "Weighted Payout", "Cum Weight", "Cum Prob", "Payout vs Cost"]
    for col, header in enumerate(headers, start=1):
        ws.cell(header_row, col, header)
        ws.cell(header_row, col).fill = HEADER_FILL
        ws.cell(header_row, col).font = HEADER_FONT
        ws.cell(header_row, col).alignment = WRAP

    for i, (sim_id, weight, raw, _, _, _, _, _) in enumerate(stats.rows):
        r = data_start + i
        ws.cell(r, 1, sim_id)
        ws.cell(r, 2, weight)
        ws.cell(r, 3, raw)
        ws.cell(r, 4, f"=C{r}/{int(PAYOUT_SCALE)}")
        ws.cell(r, 5, f"=B{r}/$B${metric_rows['Total Weight']}")
        ws.cell(r, 6, f"=D{r}*E{r}")
        ws.cell(r, 7, f"=SUM($B${data_start}:B{r})")
        ws.cell(r, 8, f"=SUM($E${data_start}:E{r})")
        ws.cell(r, 9, f"=D{r}-$B${metric_rows['Cost']}")

    for r in range(data_start, data_end + 1):
        ws.cell(r, 4).number_format = "0.000000"
        ws.cell(r, 5).number_format = "0.000000%"
        ws.cell(r, 6).number_format = "0.000000"
        ws.cell(r, 7).number_format = "0"
        ws.cell(r, 8).number_format = "0.000000%"
        ws.cell(r, 9).number_format = "0.000000"

    ws.cell(metric_rows["Cost"], 2, f"={stats.cost}")
    ws.cell(metric_rows["Total Weight"], 2, f"=SUM(B{data_start}:B{data_end})")
    ws.cell(metric_rows["Weighted Mean Payout"], 2, f"=SUM(F{data_start}:F{data_end})")
    ws.cell(metric_rows["RTP"], 2, f"=B{metric_rows['Weighted Mean Payout']}/B{metric_rows['Cost']}")
    ws.cell(metric_rows["Hit Rate"], 2, f"=SUMIF(D{data_start}:D{data_end},\">0\",E{data_start}:E{data_end})")
    ws.cell(metric_rows["Zero Rate"], 2, f"=SUMIF(D{data_start}:D{data_end},0,E{data_start}:E{data_end})")
    ws.cell(metric_rows["Below-Cost Rate"], 2, f"=SUMIFS(E{data_start}:E{data_end},D{data_start}:D{data_end},\">0\",D{data_start}:D{data_end},\"<\"&$B${metric_rows['Cost']})")
    ws.cell(metric_rows["Profit Rate"], 2, f"=SUMIFS(E{data_start}:E{data_end},D{data_start}:D{data_end},\">\"&$B${metric_rows['Cost']})")
    ws.cell(metric_rows["Variance"], 2, f"=SUMPRODUCT(E{data_start}:E{data_end},(D{data_start}:D{data_end}-$B${metric_rows['Weighted Mean Payout']})*(D{data_start}:D{data_end}-$B${metric_rows['Weighted Mean Payout']}))")
    ws.cell(metric_rows["Std Dev"], 2, f"=SQRT(B{metric_rows['Variance']})")
    ws.cell(metric_rows["Min Win"], 2, f"=MIN(D{data_start}:D{data_end})")
    ws.cell(metric_rows["Max Win"], 2, f"=MAX(D{data_start}:D{data_end})")
    ws.cell(metric_rows["Median"], 2, f"=INDEX(D{data_start}:D{data_end},COUNTIF(H{data_start}:H{data_end},\"<0.5\")+1)")
    ws.cell(metric_rows["P95"], 2, f"=INDEX(D{data_start}:D{data_end},COUNTIF(H{data_start}:H{data_end},\"<0.95\")+1)")
    ws.cell(metric_rows["P99"], 2, f"=INDEX(D{data_start}:D{data_end},COUNTIF(H{data_start}:H{data_end},\"<0.99\")+1)")

    for row_num in range(6, 21):
        ws.cell(row_num, 2).number_format = "0.000000"
    ws.cell(metric_rows["Cost"], 2).number_format = "0.00"
    for lbl in ["Hit Rate", "Zero Rate", "Below-Cost Rate", "Profit Rate"]:
        ws.cell(metric_rows[lbl], 2).number_format = "0.000000%"

    ws.freeze_panes = f"A{data_start}"
    ws.auto_filter.ref = f"A{header_row}:I{data_end}"
    set_col_widths(ws, {1: 12, 2: 12, 3: 12, 4: 12, 5: 12, 6: 15, 7: 12, 8: 12, 9: 14})


def write_json_sheet(wb: Workbook, sheet_name: str, title: str, path: Path) -> None:
    ws = wb.create_sheet(title=safe_sheet_name(sheet_name))
    data = load_json(path)
    add_title(ws, title, path.name)
    row = 3
    row = section(ws, row, "JSON Overview")
    overview = [
        ("Type", type(data).__name__),
        ("Top-Level Keys", ", ".join(data.keys()) if isinstance(data, dict) else "list"),
    ]
    row = kv_table(ws, row, overview)
    row += 1
    row = section(ws, row, "Serialized JSON")
    serialized = json.dumps(data, indent=2, ensure_ascii=False)
    ws.cell(row, 1, "Line")
    ws.cell(row, 2, "Content")
    ws.cell(row, 1).fill = HEADER_FILL
    ws.cell(row, 2).fill = HEADER_FILL
    ws.cell(row, 1).font = HEADER_FONT
    ws.cell(row, 2).font = HEADER_FONT
    for idx, line in enumerate(serialized.splitlines(), start=row + 1):
        ws.cell(idx, 1, idx - row)
        ws.cell(idx, 2, line)
        ws.cell(idx, 2).alignment = WRAP
    ws.freeze_panes = f"A{row + 1}"
    set_col_widths(ws, {1: 10, 2: 140})


def write_md_sheet(wb: Workbook, filename: str, audit_items: List[Tuple[str, Any]], lookup_map: Dict[str, LookupStats]) -> None:
    path = DOCS_DIR / filename
    ws = wb.create_sheet(title=safe_sheet_name(path.stem))
    add_title(ws, path.stem.replace("_", " "), path.name)
    row = 3
    if audit_items:
        row = section(ws, row, "Audit Metrics")
        row = kv_table(ws, row, audit_items)
        row += 1

    # Add special computed tables for the docs that contain the real math.
    if filename == "Treasure_Dice_Approved_Regular_Routes_Paytable.md":
        row = section(ws, row, "Computed Route Audit Table")
        headers = ["Mode", "Win Rate", "Multiplier", "Payout @ B=1", "Net Profit", "RTP", "Variance", "Std Dev", "Max Win"]
        table_rows = []
        for mode in ["safe_shore", "hidden_bay", "coral_reef", "storm_route", "skull_island", "kraken_waters", "lost_treasure"]:
            st = lookup_map[mode]
            table_rows.append([mode, st.hit_rate, st.max_win, st.max_win, st.max_win - 1.0, st.rtp, st.variance, st.std, st.max_win])
        data_table(ws, row, headers, table_rows)
        end_row = row + len(table_rows)
        for r in range(row + 1, end_row + 1):
            for c in range(2, 10):
                ws.cell(r, c).number_format = "0.000000"
        row = end_row + 2
    elif filename == "Treasure_Dice_Approved_Bonus_Buy_Paytable.md":
        row = section(ws, row, "Computed Bonus Outcome Distribution")
        headers = ["Payout Norm", "Weight", "Probability", "Expected Contribution", "Cumulative Probability", "Zero?", "Below Cost?", "Profit?"]
        unique = unique_outcomes(load_lookup_rows(PUBLISH_DIR / "lookUpTable_treasure_hunt_buy_0.csv"))
        table_rows = []
        for item in unique:
            table_rows.append([
                item.payout,
                item.weight,
                item.probability,
                item.weighted_contribution,
                item.cumulative_probability,
                "Yes" if item.payout == 0 else "No",
                "Yes" if 0 < item.payout < 100.0 else "No",
                "Yes" if item.payout > 100.0 else "No",
            ])
        data_table(ws, row, headers, table_rows)
        end_row = row + len(table_rows)
        for r in range(row + 1, end_row + 1):
            for c in range(1, 6):
                ws.cell(r, c).number_format = "0.000000"
        row = end_row + 2
    elif filename == "Treasure_Dice_Final_Math_Report.md":
        row = section(ws, row, "Computed Final Mode Summary")
        headers = ["Mode", "Cost", "RTP", "Hit Rate", "Std Dev", "P95", "P99", "Max Win", "Status"]
        table_rows = []
        for mode in ["safe_shore", "hidden_bay", "coral_reef", "storm_route", "skull_island", "kraken_waters", "lost_treasure", "treasure_hunt_buy"]:
            st = lookup_map[mode]
            table_rows.append([mode, st.cost, st.rtp, st.hit_rate, st.std, st.p95, st.p99, st.max_win, "Pass" if abs(st.rtp - 0.96) <= 0.01 else "Fail"])
        data_table(ws, row, headers, table_rows)
        end_row = row + len(table_rows)
        for r in range(row + 1, end_row + 1):
            for c in range(2, 8):
                ws.cell(r, c).number_format = "0.000000"
        row = end_row + 2
    elif filename == "Treasure_Dice_Phase_5_Final_Validation_Summary.md":
        row = section(ws, row, "Computed Validation Table")
        headers = ["Mode", "Cost", "Theoretical RTP", "Recalculated RTP", "Drift", "Pass"]
        validation_rows = []
        validation_data = {m["mode_id"]: m for m in load_json(DOCS_DIR / "treasure_dice_phase5_validation_summary.json")["modes"]}
        for mode in ["safe_shore", "hidden_bay", "coral_reef", "storm_route", "skull_island", "kraken_waters", "lost_treasure", "treasure_hunt_buy"]:
            entry = validation_data[mode]
            validation_rows.append([
                mode,
                entry["cost"],
                entry["theoretical_rtp"],
                entry["simulated_rtp"],
                entry["rtp_drift"],
                "Pass" if entry.get("rtp_tolerance_pass", False) else "Fail",
            ])
        data_table(ws, row, headers, validation_rows)
        end_row = row + len(validation_rows)
        for r in range(row + 1, end_row + 1):
            for c in range(2, 5):
                ws.cell(r, c).number_format = "0.000000"
            ws.cell(r, 5).number_format = "0.000000"
        row = end_row + 2

    row = section(ws, row, "Source Text")
    ws.cell(row, 1, "Line")
    ws.cell(row, 2, "Content")
    ws.cell(row, 1).fill = HEADER_FILL
    ws.cell(row, 2).fill = HEADER_FILL
    ws.cell(row, 1).font = HEADER_FONT
    ws.cell(row, 2).font = HEADER_FONT
    for idx, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=row + 1):
        ws.cell(idx, 1, idx - row)
        ws.cell(idx, 2, line)
        ws.cell(idx, 2).alignment = WRAP
    ws.freeze_panes = f"A{row + 1}"
    set_col_widths(ws, {1: 10, 2: 140})


def build_summary_sheet(ws, index: dict, lookup_map: Dict[str, LookupStats], validation: dict) -> None:
    add_title(ws, "Treasure Dice Full Audit Workbook", "Clean workbook for treasure_dice_final")
    row = 3
    row = section(ws, row, "Workbook Scope")
    scope = [
        ("Package", "math-sdk/games/treasure_dice_final"),
        ("Modes", len(index.get("modes", []))),
        ("Target RTP", 0.96),
        ("Tolerance", 0.01),
        ("All RTP Pass", validation.get("checks", {}).get("all_rtp_tolerance_pass", False)),
        ("All Weight Pass", validation.get("checks", {}).get("all_weight_checks_pass", False)),
        ("All Lookup Pass", validation.get("checks", {}).get("all_lookup_checks_pass", False)),
        ("All MaxWin Pass", validation.get("checks", {}).get("all_max_win_checks_pass", False)),
        ("All Replay Pass", validation.get("checks", {}).get("all_replay_checks_pass", False)),
    ]
    row = kv_table(ws, row, scope, {"Target RTP": "0.00%", "Tolerance": "0.00%"})
    row += 1
    row = section(ws, row, "Mode Audit Table")
    headers = ["Mode", "Cost", "RTP", "Hit Rate", "Std Dev", "P95", "P99", "Max Win", "Pass", "Sheet"]
    data_rows = []
    for mode in index["modes"]:
        sheet_name = safe_sheet_name(f"LUT_{mode['name']}")
        row_idx = row + 1 + len(data_rows)
        data_rows.append([
            mode["name"],
            f"={sheet_name}!B6",
            f"={sheet_name}!B9",
            f"={sheet_name}!B10",
            f"={sheet_name}!B15",
            f"={sheet_name}!B19",
            f"={sheet_name}!B20",
            f"={sheet_name}!B17",
            f'=IF(ABS(C{row_idx}-0.96)<=0.01,"Pass","Fail")',
            f"LUT_{mode['name']}",
        ])
    data_table(ws, row, headers, data_rows)
    end_row = row + len(data_rows)
    for r in range(row + 1, end_row + 1):
        for c in range(2, 9):
            ws.cell(r, c).number_format = "0.000000"
    set_col_widths(ws, {1: 18, 2: 10, 3: 10, 4: 10, 5: 10, 6: 10, 7: 10, 8: 10, 9: 8, 10: 20})
    ws.freeze_panes = f"A{row + 1}"


def main() -> int:
    index = load_json(PUBLISH_DIR / "index.json")
    validation = load_json(DOCS_DIR / "treasure_dice_phase5_validation_summary.json")

    lookup_map: Dict[str, LookupStats] = {}
    lookup_rows: Dict[str, List[Tuple[int, int, float]]] = {}
    for mode in index["modes"]:
        rows = load_lookup_rows(PUBLISH_DIR / mode["weights"])
        lookup_rows[mode["name"]] = rows
        lookup_map[mode["name"]] = calc_lookup_stats(mode["name"], float(mode["cost"]), rows)

    wb = Workbook()
    if hasattr(wb, "calculation"):
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    active_sheet = wb.active
    if active_sheet is not None:
        wb.remove(active_sheet)

    build_summary_sheet(wb.create_sheet("Summary"), index, lookup_map, validation)
    write_json_sheet(wb, "Index_JSON", "Index JSON Audit", PUBLISH_DIR / "index.json")
    write_json_sheet(wb, "Config_JSON", "Config JSON Audit", CONFIG_DIR / "config.json")
    write_json_sheet(wb, "FE_Config_JSON", "Frontend Config JSON Audit", CONFIG_DIR / "config_fe_treasure_dice.json")
    write_json_sheet(wb, "Mapping_JSON", "Event Mapping JSON Audit", MAPPING_DIR / "treasure_dice_mapping_spec.json")
    write_json_sheet(wb, "Validation_JSON", "Validation Summary JSON Audit", DOCS_DIR / "treasure_dice_phase5_validation_summary.json")

    for mode in index["modes"]:
        ws = wb.create_sheet(safe_sheet_name(f"LUT_{mode['name']}"))
        write_lookup_sheet(ws, lookup_map[mode["name"]])

    # Add a bonus distribution sheet because treasure_hunt_buy has the richest audit surface.
    bonus_ws = wb.create_sheet("Bonus_Outcome_Audit")
    bonus_stats = lookup_map["treasure_hunt_buy"]
    bonus_unique = unique_outcomes(lookup_rows["treasure_hunt_buy"])
    add_title(bonus_ws, "Bonus Outcome Audit - treasure_hunt_buy", "Unique payout distribution aggregated from the lookup table")
    bonus_ws.cell(4, 1, "Key")
    bonus_ws.cell(4, 2, "Value")
    bonus_ws.cell(4, 1).fill = HEADER_FILL
    bonus_ws.cell(4, 2).fill = HEADER_FILL
    bonus_ws.cell(4, 1).font = HEADER_FONT
    bonus_ws.cell(4, 2).font = HEADER_FONT

    metric_labels = [
        "Cost",
        "Total Weight",
        "Weighted Mean Payout",
        "RTP",
        "Hit Rate",
        "Zero Rate",
        "Below-Cost Rate",
        "Profit Rate",
        "Variance",
        "Std Dev",
        "Min Win",
        "Max Win",
        "Median",
        "P95",
        "P99",
    ]
    metric_rows = {label: 6 + idx for idx, label in enumerate(metric_labels)}
    for label, row_num in metric_rows.items():
        bonus_ws.cell(row_num, 1, label)
        bonus_ws.cell(row_num, 1).font = BOLD_FONT

    header_row = 22
    data_start = 23
    data_end = data_start + len(bonus_unique) - 1
    headers = ["Payout Norm", "Weight", "Probability", "Expected Contribution", "Cumulative Probability", "Zero?", "Below Cost?", "Profit?"]
    for col, header in enumerate(headers, start=1):
        bonus_ws.cell(header_row, col, header)
        bonus_ws.cell(header_row, col).fill = HEADER_FILL
        bonus_ws.cell(header_row, col).font = HEADER_FONT
        bonus_ws.cell(header_row, col).alignment = WRAP

    for i, item in enumerate(bonus_unique):
        r = data_start + i
        bonus_ws.cell(r, 1, item.payout)
        bonus_ws.cell(r, 2, item.weight)
        bonus_ws.cell(r, 3, f"=B{r}/$B${metric_rows['Total Weight']}")
        bonus_ws.cell(r, 4, f"=A{r}*C{r}")
        bonus_ws.cell(r, 5, f"=SUM($C${data_start}:C{r})")
        bonus_ws.cell(r, 6, f'=IF(A{r}=0,"Yes","No")')
        bonus_ws.cell(r, 7, f'=IF(AND(A{r}>0,A{r}<100),"Yes","No")')
        bonus_ws.cell(r, 8, f'=IF(A{r}>100,"Yes","No")')
        bonus_ws.cell(r, 1).number_format = "0.000000"
        bonus_ws.cell(r, 3).number_format = "0.000000%"
        bonus_ws.cell(r, 4).number_format = "0.000000"
        bonus_ws.cell(r, 5).number_format = "0.000000%"

    bonus_ws.cell(metric_rows["Cost"], 2, "=100")
    bonus_ws.cell(metric_rows["Total Weight"], 2, f"=SUM(B{data_start}:B{data_end})")
    bonus_ws.cell(metric_rows["Weighted Mean Payout"], 2, f"=SUM(D{data_start}:D{data_end})")
    bonus_ws.cell(metric_rows["RTP"], 2, f"=B{metric_rows['Weighted Mean Payout']}/B{metric_rows['Cost']}")
    bonus_ws.cell(metric_rows["Hit Rate"], 2, f"=SUMIF(A{data_start}:A{data_end},\">0\",C{data_start}:C{data_end})")
    bonus_ws.cell(metric_rows["Zero Rate"], 2, f"=SUMIF(A{data_start}:A{data_end},0,C{data_start}:C{data_end})")
    bonus_ws.cell(metric_rows["Below-Cost Rate"], 2, f"=SUMIFS(C{data_start}:C{data_end},A{data_start}:A{data_end},\">0\",A{data_start}:A{data_end},\"<100\")")
    bonus_ws.cell(metric_rows["Profit Rate"], 2, f"=SUMIFS(C{data_start}:C{data_end},A{data_start}:A{data_end},\">100\")")
    bonus_ws.cell(metric_rows["Variance"], 2, f"=SUMPRODUCT(C{data_start}:C{data_end},(A{data_start}:A{data_end}-$B${metric_rows['Weighted Mean Payout']})*(A{data_start}:A{data_end}-$B${metric_rows['Weighted Mean Payout']}))")
    bonus_ws.cell(metric_rows["Std Dev"], 2, f"=SQRT(B{metric_rows['Variance']})")
    bonus_ws.cell(metric_rows["Min Win"], 2, f"=MIN(A{data_start}:A{data_end})")
    bonus_ws.cell(metric_rows["Max Win"], 2, f"=MAX(A{data_start}:A{data_end})")
    bonus_ws.cell(metric_rows["Median"], 2, f"=INDEX(A{data_start}:A{data_end},COUNTIF(E{data_start}:E{data_end},\"<0.5\")+1)")
    bonus_ws.cell(metric_rows["P95"], 2, f"=INDEX(A{data_start}:A{data_end},COUNTIF(E{data_start}:E{data_end},\"<0.95\")+1)")
    bonus_ws.cell(metric_rows["P99"], 2, f"=INDEX(A{data_start}:A{data_end},COUNTIF(E{data_start}:E{data_end},\"<0.99\")+1)")

    for row_num in range(6, 21):
        bonus_ws.cell(row_num, 2).number_format = "0.000000"
    bonus_ws.cell(metric_rows["Cost"], 2).number_format = "0.00"
    for lbl in ["Hit Rate", "Zero Rate", "Below-Cost Rate", "Profit Rate"]:
        bonus_ws.cell(metric_rows[lbl], 2).number_format = "0.000000%"

    bonus_ws.freeze_panes = f"A{data_start}"
    bonus_ws.auto_filter.ref = f"A{header_row}:H{data_end}"
    set_col_widths(bonus_ws, {1: 14, 2: 12, 3: 12, 4: 18, 5: 20, 6: 10, 7: 12, 8: 10})

    regular_audits = [
        ("Treasure_Dice_Final_Math_Report.md", [("Model Version", "treasure_dice_release_math_v1_2026_06_04"), ("Modes", len(index["modes"])), ("All RTP Pass", validation["checks"]["all_rtp_tolerance_pass"])]),
        ("Treasure_Dice_Approved_Regular_Routes_Paytable.md", [("Source", "Lookup tables + approved route math"), ("Regular Modes", 7), ("Target RTP", 0.96), ("Tolerance", 0.01)]),
        ("Treasure_Dice_Approved_Bonus_Buy_Paytable.md", [("Mode", "treasure_hunt_buy"), ("Cost", 100.0), ("Target RTP", 0.96), ("RTP", bonus_stats.rtp), ("Hit Rate", bonus_stats.hit_rate), ("Max Win", bonus_stats.max_win)]),
        ("Treasure_Dice_Phase_3_Event_Mapping_Report.md", [("Mapping File", "treasure_dice_phase3_mapping_spec.json"), ("Deterministic Replay", True), ("Presentation Only", True)]),
        ("Treasure_Dice_Phase_5_Simulation_Report.md", [("Purpose", "Simulation evidence"), ("Modes", len(index["modes"]))]),
        ("Treasure_Dice_Phase_5_Volatility_Report.md", [("Regular Std Range", f"{min(lookup_map[m].std for m in lookup_map if m != 'treasure_hunt_buy'):.6f} - {max(lookup_map[m].std for m in lookup_map if m != 'treasure_hunt_buy'):.6f}"), ("Bonus Std", bonus_stats.std), ("Bonus Variance", bonus_stats.variance)]),
        ("Treasure_Dice_Phase_5_Final_Validation_Summary.md", [("All RTP Pass", validation["checks"].get("all_rtp_tolerance_pass", False)), ("All Weight Pass", validation["checks"].get("all_weight_checks_pass", False)), ("All Lookup Pass", validation["checks"].get("all_lookup_checks_pass", False)), ("All MaxWin Pass", validation["checks"].get("all_max_win_checks_pass", False)), ("All Replay Pass", validation["checks"].get("all_replay_checks_pass", False))]),
        ("Treasure_Dice_Phase_6_Precision_Rounding_And_Frontend_Contract.md", [("Policy", "Stake Engine integer money with six decimal places"), ("Small Bets", ", ".join(str(x) for x in SMALL_BETS)), ("RTP Retention", "Exact under final settlement model")]),
        ("Treasure_Dice_Exposure_Table_Pending_Operator_Cap.md", [("Status", "Pending operator payout cap"), ("Formula", "max allowed base bet = payout cap / max win multiplier"), ("Highest Regular Max Win", lookup_map["lost_treasure"].max_win), ("Bonus Max Win", bonus_stats.max_win)]),
        ("Treasure_Dice_Phase_6_Risk_Assessment.md", [("Package Status", "Ready"), ("Acceptance Status", "Pass"), ("Key Risk", "Operator payout cap not yet supplied")]),
        ("Treasure_Dice_Phase_6_Acceptance_Checklist.md", [("RTP Checks", "Pass"), ("Replay Checks", "Pass"), ("Lookup Checks", "Pass"), ("Max Win Checks", "Pass")]),
        ("Treasure_Dice_Phase_6_Publication_Status_Note.md", [("Publication Status", "Ready"), ("Acceptance Status", "Pass"), ("Package", "treasure_dice_final")]),
    ]
    for fname, audit in regular_audits:
        write_md_sheet(wb, fname, audit, lookup_map)

    temp_output = OUTPUT.with_suffix(".tmp.xlsx")
    wb.save(temp_output)
    with zipfile.ZipFile(temp_output, "r") as archive:
        bad = archive.testzip()
        if bad is not None:
            raise RuntimeError(f"Workbook archive failed validation: {bad}")
        if "xl/workbook.xml" not in archive.namelist():
            raise RuntimeError("Workbook archive missing xl/workbook.xml")
    temp_output.replace(OUTPUT)
    check = load_workbook(OUTPUT, read_only=True)
    print(f"Wrote {OUTPUT}")
    print(f"Sheets: {len(check.sheetnames)}")
    print("Sheet names:", check.sheetnames)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
