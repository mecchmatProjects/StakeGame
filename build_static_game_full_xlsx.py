from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FILL = PatternFill("solid", fgColor="B4C6E7")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")

TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)
HEADER_FONT = Font(bold=True)
BOLD_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


@dataclass
class ModeStats:
    mode: str
    cost: float
    rows_configured: int | None
    rtp_target: float | None
    expected_buckets: int | None
    outcome_rows: int
    unique_multipliers: int
    bucket_structure_pass: bool | None
    total_weight: int
    weighted_mean: float
    rtp_actual: float
    hit_rate: float
    zero_rate: float
    profit_rate: float
    variance: float
    std: float
    min_win: float
    max_win: float
    max_win_probability: float
    max_exposure: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a generic static-game Excel workbook from package artifacts.")
    parser.add_argument("--root", required=True, help="Path to game package root, for example math-sdk/games/azteck_plinko_final")
    parser.add_argument("--output", default="", help="Optional output xlsx path. Defaults to <root>/<WorkbookName>.xlsx")
    parser.add_argument("--name", default="StaticGameFull", help="Workbook base name when --output is omitted")
    parser.add_argument("--target-rtp", type=float, default=None, help="Fallback target RTP if not present in index")
    parser.add_argument("--rtp-tol", type=float, default=0.01, help="Absolute RTP tolerance for pass or fail checks")
    parser.add_argument("--payout-scale", type=float, default=100.0, help="Divide lookup payout values by this scale")
    parser.add_argument("--max-bet", type=float, default=500.0, help="Max bet assumption used for Max Exposure reporting")
    return parser.parse_args()


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def add_title(ws, title: str, subtitle: str = "") -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = WRAP
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True)
        ws["A2"].alignment = WRAP


def section(ws, row: int, title: str) -> int:
    ws.cell(row, 1, title)
    ws.cell(row, 1).font = BOLD_FONT
    ws.cell(row, 1).fill = SECTION_FILL
    return row + 1


def set_col_widths(ws, widths: dict[int, float]) -> None:
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width


def kv_table(ws, start_row: int, items: list[tuple[str, Any]]) -> int:
    ws.cell(start_row, 1, "Key")
    ws.cell(start_row, 2, "Value")
    ws.cell(start_row, 1).fill = HEADER_FILL
    ws.cell(start_row, 2).fill = HEADER_FILL
    ws.cell(start_row, 1).font = HEADER_FONT
    ws.cell(start_row, 2).font = HEADER_FONT
    row = start_row + 1
    for key, value in items:
        ws.cell(row, 1, key)
        ws.cell(row, 2, value)
        ws.cell(row, 1).alignment = TOP
        ws.cell(row, 2).alignment = TOP
        row += 1
    return row


def data_table(ws, start_row: int, headers: list[str], rows: list[list[Any]]) -> int:
    for col, header in enumerate(headers, start=1):
        ws.cell(start_row, col, header)
        ws.cell(start_row, col).fill = HEADER_FILL
        ws.cell(start_row, col).font = HEADER_FONT
        ws.cell(start_row, col).alignment = WRAP

    row = start_row + 1
    for record in rows:
        for col, value in enumerate(record, start=1):
            ws.cell(row, col, value)
            ws.cell(row, col).alignment = TOP
        row += 1
    return row


def read_lookup_rows(path: Path) -> list[tuple[int, int, float]]:
    out: list[tuple[int, int, float]] = []
    with path.open("r", encoding="utf-8", newline="") as handle:
        for row in csv.reader(handle):
            if len(row) < 3:
                continue
            try:
                sim_id = int(float(row[0]))
                weight = int(float(row[1]))
                payout = float(row[2])
            except ValueError:
                continue
            out.append((sim_id, weight, payout))
    return out


def detect_mode_costs(index_data: dict[str, Any]) -> dict[str, float]:
    costs: dict[str, float] = {}
    modes = index_data.get("modes", [])
    if isinstance(modes, list):
        for mode in modes:
            if not isinstance(mode, dict):
                continue
            name = mode.get("name") or mode.get("mode")
            if not isinstance(name, str):
                continue
            cost = mode.get("cost")
            if isinstance(cost, (int, float)):
                costs[name] = float(cost)
    return costs


def detect_mode_rows(index_data: dict[str, Any]) -> dict[str, int]:
    rows_map: dict[str, int] = {}
    modes = index_data.get("modes", [])
    if isinstance(modes, list):
        for mode in modes:
            if not isinstance(mode, dict):
                continue
            name = mode.get("name") or mode.get("mode")
            rows = mode.get("rows")
            if isinstance(name, str) and isinstance(rows, int):
                rows_map[name] = rows
    return rows_map


def detect_target_rtp(index_data: dict[str, Any], mode: str, fallback: float | None) -> float | None:
    modes = index_data.get("modes", [])
    if isinstance(modes, list):
        for item in modes:
            if not isinstance(item, dict):
                continue
            name = item.get("name") or item.get("mode")
            if name != mode:
                continue
            for key in ("rtp_target", "target_rtp", "rtp"):
                value = item.get(key)
                if isinstance(value, (int, float)):
                    return float(value)
    for key in ("rtp_target", "target_rtp", "rtp"):
        value = index_data.get(key)
        if isinstance(value, (int, float)):
            return float(value)
    return fallback


def compute_mode_stats(
    mode: str,
    cost: float,
    rows_configured: int | None,
    rtp_target: float | None,
    rows: list[tuple[int, int, float]],
    payout_scale: float,
) -> ModeStats:
    outcome_rows = len(rows)
    expected_buckets = (rows_configured + 1) if rows_configured is not None else None
    bucket_structure_pass = (outcome_rows == expected_buckets) if expected_buckets is not None else None

    total_weight = sum(weight for _, weight, _ in rows)
    if total_weight <= 0:
        return ModeStats(
            mode=mode,
            cost=cost,
            rows_configured=rows_configured,
            rtp_target=rtp_target,
            expected_buckets=expected_buckets,
            outcome_rows=outcome_rows,
            unique_multipliers=0,
            bucket_structure_pass=bucket_structure_pass,
            total_weight=0,
            weighted_mean=0.0,
            rtp_actual=0.0,
            hit_rate=0.0,
            zero_rate=0.0,
            profit_rate=0.0,
            variance=0.0,
            std=0.0,
            min_win=0.0,
            max_win=0.0,
            max_win_probability=0.0,
            max_exposure=0.0,
        )

    payouts = [raw / payout_scale for _, _, raw in rows]
    probs = [weight / total_weight for _, weight, _ in rows]

    weighted_mean = sum(p * x for p, x in zip(probs, payouts))
    rtp_actual = (weighted_mean / cost) if cost > 0 else 0.0
    hit_rate = sum(p for p, x in zip(probs, payouts) if x > 0)
    zero_rate = sum(p for p, x in zip(probs, payouts) if x == 0)
    profit_rate = sum(p for p, x in zip(probs, payouts) if x > cost)
    variance = sum(p * ((x - weighted_mean) ** 2) for p, x in zip(probs, payouts))
    std = math.sqrt(variance)
    max_win = max(payouts)
    max_win_probability = sum(p for p, x in zip(probs, payouts) if abs(x - max_win) <= 1e-12)
    unique_multipliers = len(set(round(x, 12) for x in payouts))

    return ModeStats(
        mode=mode,
        cost=cost,
        rows_configured=rows_configured,
        rtp_target=rtp_target,
        expected_buckets=expected_buckets,
        outcome_rows=outcome_rows,
        unique_multipliers=unique_multipliers,
        bucket_structure_pass=bucket_structure_pass,
        total_weight=total_weight,
        weighted_mean=weighted_mean,
        rtp_actual=rtp_actual,
        hit_rate=hit_rate,
        zero_rate=zero_rate,
        profit_rate=profit_rate,
        variance=variance,
        std=std,
        min_win=min(payouts),
        max_win=max_win,
        max_win_probability=max_win_probability,
        max_exposure=max_win,
    )


def test_vectors_for_rows(rows: list[tuple[int, int, float]]) -> list[tuple[int, int, float, int, int]]:
    if not rows:
        return []

    intervals: list[tuple[int, int, int, float]] = []
    running = 0
    for sim_id, weight, payout_raw in rows:
        if weight <= 0:
            continue
        start = running
        end = running + weight - 1
        intervals.append((sim_id, start, end, payout_raw))
        running += weight

    if not intervals:
        return []

    total = running
    draws = [0, max(0, total // 4), max(0, total // 2), max(0, (3 * total) // 4), max(0, total - 1)]
    vectors: list[tuple[int, int, float, int, int]] = []

    for draw in draws:
        for sim_id, start, end, payout_raw in intervals:
            if start <= draw <= end:
                vectors.append((draw, sim_id, payout_raw, start, end))
                break

    return vectors


def main() -> None:
    args = parse_args()
    root = Path(args.root).resolve()
    publish_dir = root / "artifacts" / "publish_files"
    docs_dir = root / "docs"

    index_path = publish_dir / "index.json"
    if not index_path.exists():
        raise FileNotFoundError(f"Missing index file: {index_path}")

    output = Path(args.output).resolve() if args.output else (root / f"{args.name}.xlsx")
    index_data = load_json(index_path)
    mode_costs = detect_mode_costs(index_data)
    mode_rows_config = detect_mode_rows(index_data)

    lookup_files = sorted(publish_dir.glob("lookup_table_*.csv"))
    if not lookup_files:
        lookup_files = sorted(publish_dir.glob("lookUpTable_*.csv"))
    if not lookup_files:
        lookup_files = sorted((root / "artifacts" / "lookup_tables").glob("*.csv"))

    has_lookup_data = bool(lookup_files)

    by_mode_rows: dict[str, list[tuple[int, int, float]]] = defaultdict(list)
    for file in lookup_files:
        stem = file.stem
        if stem.startswith("lookup_table_"):
            mode = stem.replace("lookup_table_", "", 1)
        elif stem.startswith("lookUpTable_"):
            mode = stem.replace("lookUpTable_", "", 1)
            if mode.endswith("_0"):
                mode = mode[:-2]
        else:
            mode = stem
        by_mode_rows[mode] = read_lookup_rows(file)

    stats: list[ModeStats] = []
    mode_targets: dict[str, float | None] = {}
    for mode, rows in sorted(by_mode_rows.items()):
        cost = mode_costs.get(mode, 1.0)
        configured_rows = mode_rows_config.get(mode)
        target = detect_target_rtp(index_data, mode, args.target_rtp)
        mode_targets[mode] = target
        stats.append(compute_mode_stats(mode, cost, configured_rows, target, rows, args.payout_scale))

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    add_title(ws, f"{args.name} Workbook", f"Generated from {root}")
    set_col_widths(ws, {1: 34, 2: 12, 3: 12, 4: 16, 5: 12, 6: 16, 7: 20, 8: 12, 9: 12, 10: 12, 11: 12, 12: 12, 13: 12, 14: 12, 15: 14, 16: 12, 17: 12, 18: 20, 19: 14})

    row = 4
    row = section(ws, row, "Package")
    row = kv_table(
        ws,
        row,
        [
            ("Root", str(root)),
            ("Index", str(index_path)),
            ("Lookup files", len(lookup_files)),
            ("Docs dir exists", docs_dir.exists()),
            ("RTP tolerance", args.rtp_tol),
            ("Payout scale", args.payout_scale),
            ("Max bet assumption", args.max_bet),
        ],
    )

    row += 1
    if not has_lookup_data:
        row = section(ws, row, "Scaffold Warning")
        row = kv_table(
            ws,
            row,
            [
                ("Status", "No lookup tables found"),
                ("Impact", "Mode metrics are placeholders until publish files are generated"),
            ],
        )
        row += 1

    row = section(ws, row, "Mode Metrics")
    mode_metrics_header_row = row
    headers = [
        "Mode",
        "Cost",
        "Rows Config",
        "Expected Buckets",
        "Rows",
        "Unique Multipliers",
        "Bucket Structure Pass",
        "RTP Target",
        "RTP Actual",
        "Abs Drift",
        "RTP Pass",
        "Hit Rate",
        "Zero Rate",
        "Profit Rate",
        "Weighted Mean",
        "Std",
        "Max Win",
        "Max Win Probability",
        "Max Exposure",
    ]

    mode_rows: list[list[Any]] = []
    all_pass = True
    for s in stats:
        drift = abs(s.rtp_actual - s.rtp_target) if s.rtp_target is not None else None
        rtp_pass = (drift <= args.rtp_tol) if drift is not None else "N/A"
        if isinstance(rtp_pass, bool):
            all_pass = all_pass and rtp_pass
        mode_rows.append(
            [
                s.mode,
                s.cost,
                s.rows_configured,
                s.expected_buckets,
                s.outcome_rows,
                s.unique_multipliers,
                s.bucket_structure_pass,
                s.rtp_target,
                s.rtp_actual,
                drift,
                rtp_pass,
                s.hit_rate,
                s.zero_rate,
                s.profit_rate,
                s.weighted_mean,
                s.std,
                s.max_win,
                s.max_win_probability,
                s.max_exposure,
            ]
        )

    mode_metrics_data_start = mode_metrics_header_row + 1
    row = data_table(ws, row, headers, mode_rows)
    mode_metrics_data_end = row - 1
    mode_summary_row: dict[str, int] = {}
    for idx, record in enumerate(mode_rows, start=mode_metrics_data_start):
        mode_summary_row[str(record[0])] = idx

    for r in range(row - len(mode_rows), row):
        pass_cell = ws.cell(r, 11)
        if pass_cell.value is True:
            pass_cell.fill = GREEN_FILL
        elif pass_cell.value is False:
            pass_cell.fill = RED_FILL

    row += 1
    row = section(ws, row, "Summary Checks")
    summary_checks_table_start = row
    row = kv_table(
        ws,
        row,
        [
            ("Modes analyzed", f"=ROWS(A{mode_metrics_data_start}:A{mode_metrics_data_end})"),
            ("RTP pass count", f"=COUNTIF(K{mode_metrics_data_start}:K{mode_metrics_data_end},TRUE)"),
            ("Probability sum pass count", "=0"),
            ("All RTP pass", f"=COUNTIF(K{mode_metrics_data_start}:K{mode_metrics_data_end},FALSE)=0"),
            ("Bucket structure pass count", f"=COUNTIF(G{mode_metrics_data_start}:G{mode_metrics_data_end},TRUE)"),
            ("Min RTP Drift", f"=MIN(J{mode_metrics_data_start}:J{mode_metrics_data_end})"),
            ("Max RTP Drift", f"=MAX(J{mode_metrics_data_start}:J{mode_metrics_data_end})"),
        ],
    )

    row += 1
    row = section(ws, row, "Validation Snapshot")
    row = kv_table(
        ws,
        row,
        [
            ("Modes analyzed", f"=ROWS(A{mode_metrics_data_start}:A{mode_metrics_data_end})"),
            ("All RTP pass", f"=COUNTIF(K{mode_metrics_data_start}:K{mode_metrics_data_end},FALSE)=0"),
            ("Generated workbook", str(output)),
        ],
    )

    ws.freeze_panes = "A5"

    detail_headers = [
        "sim_id",
        "weight",
        "payout_raw",
        "payout_scaled",
        "probability",
        "weighted_payout",
        "cumulative_probability",
    ]

    audit_refs: dict[str, dict[str, str]] = {}

    for mode, rows_for_mode in sorted(by_mode_rows.items()):
        sheet = wb.create_sheet(title=(mode[:31] if mode else "mode"))
        add_title(sheet, f"{mode} Lookup Detail", "Derived from lookup rows")
        set_col_widths(sheet, {1: 12, 2: 12, 3: 14, 4: 16, 5: 16, 6: 18, 7: 22, 8: 14})

        mode_cost = mode_costs.get(mode, 1.0)
        mode_target = mode_targets.get(mode, args.target_rtp)
        sheet["A3"] = "Cost"
        sheet["B3"] = mode_cost
        sheet["C3"] = "Target RTP"
        sheet["D3"] = mode_target if mode_target is not None else ""
        sheet["E3"] = "RTP Tol"
        sheet["F3"] = args.rtp_tol
        sheet["G3"] = "Payout Scale"
        sheet["H3"] = args.payout_scale

        r = 5
        r = section(sheet, r, "Rows")

        for col, header in enumerate(detail_headers, start=1):
            sheet.cell(r, col, header)
            sheet.cell(r, col).fill = HEADER_FILL
            sheet.cell(r, col).font = HEADER_FONT
            sheet.cell(r, col).alignment = WRAP

        data_start = r + 1
        data_end = data_start + len(rows_for_mode) - 1

        total_row = data_end + 2
        sheet.cell(total_row, 1, "CHECK_SUMS")
        sheet.cell(total_row, 1).font = BOLD_FONT
        sheet.cell(total_row, 1).fill = SECTION_FILL

        # Totals and checks (formula-driven for audit use).
        sheet.cell(total_row, 2, f"=SUM(B{data_start}:B{data_end})")
        sheet.cell(total_row, 3, "payout_raw_sum")
        sheet.cell(total_row, 4, f"=SUM(D{data_start}:D{data_end})")
        sheet.cell(total_row, 5, f"=SUM(E{data_start}:E{data_end})")
        sheet.cell(total_row, 6, f"=SUM(F{data_start}:F{data_end})")
        sheet.cell(total_row, 7, f"=IF($B$3>0,F{total_row}/$B$3,0)")
        sheet.cell(total_row, 8, f"=IF($D$3=\"\",\"N/A\",ABS(G{total_row}-$D$3)<=$F$3)")

        for idx, (sim_id, weight, raw) in enumerate(rows_for_mode, start=data_start):
            sheet.cell(idx, 1, sim_id)
            sheet.cell(idx, 2, weight)
            sheet.cell(idx, 3, raw)
            sheet.cell(idx, 4, f"=C{idx}/$H$3")
            sheet.cell(idx, 5, f"=IF($B${total_row}>0,B{idx}/$B${total_row},0)")
            sheet.cell(idx, 6, f"=D{idx}*E{idx}")
            if idx == data_start:
                sheet.cell(idx, 7, f"=E{idx}")
            else:
                sheet.cell(idx, 7, f"=G{idx-1}+E{idx}")

        sheet.cell(total_row + 1, 1, "rtp_drift")
        sheet.cell(total_row + 1, 2, f"=IF($D$3=\"\",\"\",G{total_row}-$D$3)")
        sheet.cell(total_row + 2, 1, "prob_sum_check")
        sheet.cell(total_row + 2, 2, f"=ABS(E{total_row}-1)<=1E-9")

        sheet.freeze_panes = "A7"

        audit_refs[mode] = {
            "sheet": sheet.title,
            "data_start": str(data_start),
            "data_end": str(data_end),
            "weight_sum": f"B{total_row}",
            "payout_sum": f"D{total_row}",
            "prob_sum": f"E{total_row}",
            "weighted_mean": f"F{total_row}",
            "rtp": f"G{total_row}",
            "pass": f"H{total_row}",
            "drift": f"B{total_row + 1}",
            "max_win": f"B{total_row + 3}",
            "max_win_prob": f"B{total_row + 4}",
            "outcome_rows": f"B{total_row + 5}",
            "unique_mult": f"B{total_row + 6}",
            "max_exposure": f"B{total_row + 7}",
        }

        sheet.cell(total_row + 3, 1, "max_win_value")
        sheet.cell(total_row + 3, 2, f"=MAX(D{data_start}:D{data_end})")
        sheet.cell(total_row + 4, 1, "max_win_probability")
        sheet.cell(total_row + 4, 2, f"=SUMIFS(E{data_start}:E{data_end},D{data_start}:D{data_end},B{total_row + 3})")
        sheet.cell(total_row + 5, 1, "outcome_rows")
        sheet.cell(total_row + 5, 2, f"=ROWS(A{data_start}:A{data_end})")
        sheet.cell(total_row + 6, 1, "unique_multipliers")
        sheet.cell(total_row + 6, 2, f"=SUMPRODUCT(1/COUNTIF(D{data_start}:D{data_end},D{data_start}:D{data_end}))")
        sheet.cell(total_row + 7, 1, "max_exposure_with_max_bet")
        sheet.cell(total_row + 7, 2, f"=B{total_row + 3}*{args.max_bet}")

    # Convert Summary Mode Metrics rows to formula-driven values from per-mode sheets.
    for mode, srow in mode_summary_row.items():
        refs = audit_refs.get(mode)
        if not refs:
            continue

        sref = f"'{refs['sheet']}'"
        dstart = refs["data_start"]
        dend = refs["data_end"]

        ws.cell(srow, 2, f"={sref}!$B$3")
        ws.cell(srow, 3, mode_rows_config.get(mode))
        ws.cell(srow, 4, f'=IF(C{srow}="","",C{srow}+1)')
        ws.cell(srow, 5, f"={sref}!{refs['outcome_rows']}")
        ws.cell(srow, 6, f"={sref}!{refs['unique_mult']}")
        ws.cell(srow, 7, f'=IF(C{srow}="","N/A",F{srow}=ROUNDUP((C{srow}+1)/2,0))')
        ws.cell(srow, 8, f"={sref}!$D$3")
        ws.cell(srow, 9, f"={sref}!{refs['rtp']}")
        ws.cell(srow, 10, f"=ABS(I{srow}-H{srow})")
        ws.cell(srow, 11, f"={sref}!{refs['pass']}")
        ws.cell(srow, 12, f'=SUMIFS({sref}!E{dstart}:E{dend},{sref}!D{dstart}:D{dend},">0")')
        ws.cell(srow, 13, f'=SUMIFS({sref}!E{dstart}:E{dend},{sref}!D{dstart}:D{dend},"=0")')
        ws.cell(srow, 14, f'=SUMIFS({sref}!E{dstart}:E{dend},{sref}!D{dstart}:D{dend},">"&B{srow})')
        ws.cell(srow, 15, f"={sref}!{refs['weighted_mean']}")
        ws.cell(srow, 16, f"=SQRT(SUMPRODUCT({sref}!E{dstart}:E{dend},({sref}!D{dstart}:D{dend}-O{srow})^2))")
        ws.cell(srow, 17, f"={sref}!{refs['max_win']}")
        ws.cell(srow, 18, f"={sref}!{refs['max_win_prob']}")
        ws.cell(srow, 19, f"={sref}!{refs['max_exposure']}")

    row += 1
    row = section(ws, row, "Formula Audit Metrics")
    audit_headers = [
        "Mode",
        "Cost",
        "Target RTP",
        "Probability Sum",
        "Payout Sum",
        "Weighted Mean",
        "RTP",
        "RTP Drift",
        "RTP Pass",
        "Weight Sum",
    ]
    audit_header_row = row
    row = data_table(ws, row, audit_headers, [])

    for mode in sorted(audit_refs):
        refs = audit_refs[mode]
        sref = f"'{refs['sheet']}'"
        ws.cell(row, 1, mode)
        ws.cell(row, 2, f"={sref}!$B$3")
        ws.cell(row, 3, f"={sref}!$D$3")
        ws.cell(row, 4, f"={sref}!{refs['prob_sum']}")
        ws.cell(row, 5, f"={sref}!{refs['payout_sum']}")
        ws.cell(row, 6, f"={sref}!{refs['weighted_mean']}")
        ws.cell(row, 7, f"={sref}!{refs['rtp']}")
        ws.cell(row, 8, f"={sref}!{refs['drift']}")
        ws.cell(row, 9, f"={sref}!{refs['pass']}")
        ws.cell(row, 10, f"={sref}!{refs['weight_sum']}")
        row += 1

    # Repair summary check formulas to point at the actual probability-sum column.
    probability_sum_check_row = summary_checks_table_start + 3
    audit_data_start = audit_header_row + 1
    audit_data_end = row - 1
    ws.cell(
        probability_sum_check_row,
        2,
        (
            f"=IF({audit_data_end}<{audit_data_start},0,"
            f"COUNTIFS(D{audit_data_start}:D{audit_data_end},\"<>\","
            f"D{audit_data_start}:D{audit_data_end},\">=0.999999999\","
            f"D{audit_data_start}:D{audit_data_end},\"<=1.000000001\"))"
        ),
    )

    vectors_ws = wb.create_sheet(title="TestVectors")
    add_title(vectors_ws, "Test Vector List", "Deterministic draw points and expected outcomes from lookup tables")
    set_col_widths(vectors_ws, {1: 28, 2: 10, 3: 12, 4: 20, 5: 16, 6: 24, 7: 14, 8: 14})

    vrow = 4
    vrow = section(vectors_ws, vrow, "Vectors")
    vector_headers = [
        "Mode",
        "Vector",
        "Draw",
        "Expected Outcome ID",
        "Expected Payout",
        "Expected Payout Normalized",
        "Interval Start",
        "Interval End",
    ]
    vrow = data_table(vectors_ws, vrow, vector_headers, [])

    vector_no = 1
    for mode in sorted(by_mode_rows):
        vectors = test_vectors_for_rows(by_mode_rows[mode])
        mode_cost = mode_costs.get(mode, 1.0)
        for draw, sim_id, payout_raw, start, end in vectors:
            payout = payout_raw / args.payout_scale
            vectors_ws.cell(vrow, 1, mode)
            vectors_ws.cell(vrow, 2, vector_no)
            vectors_ws.cell(vrow, 3, draw)
            vectors_ws.cell(vrow, 4, sim_id)
            vectors_ws.cell(vrow, 5, payout)
            vectors_ws.cell(vrow, 6, (payout / mode_cost) if mode_cost > 0 else 0.0)
            vectors_ws.cell(vrow, 7, start)
            vectors_ws.cell(vrow, 8, end)
            vrow += 1
            vector_no += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"Workbook generated: {output}")


if __name__ == "__main__":
    main()
