from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


ROOT = Path(__file__).resolve().parent
OUT_DOCX = ROOT / "Treasure_Dice_Mathematician_Artifacts_Package.docx"
VALIDATION_JSON = ROOT / "treasure_dice_phase5_validation_summary.json"
PUBLISH_DIR = ROOT / "math-sdk" / "games" / "treasure_dice" / "library" / "publish_files"
CONFIG_DIR = ROOT / "math-sdk" / "games" / "treasure_dice" / "library" / "configs"
LOOKUP_DIR = ROOT / "math-sdk" / "games" / "treasure_dice" / "library" / "lookup_tables"
FORCE_DIR = ROOT / "math-sdk" / "games" / "treasure_dice" / "library" / "forces"


def rel(path: Path) -> str:
    return str(path.relative_to(ROOT)).replace("\\", "/")


def add_file_list(doc: Document, heading: str, folder: Path) -> None:
    doc.add_heading(heading, level=2)
    if not folder.exists():
        doc.add_paragraph(f"Missing folder: {rel(folder)}")
        return
    files = sorted([p for p in folder.iterdir() if p.is_file()])
    if not files:
        doc.add_paragraph(f"No files found in {rel(folder)}")
        return
    for file_path in files:
        doc.add_paragraph(rel(file_path), style="List Bullet")


def add_deliverables_table(doc: Document, validation_data: dict) -> None:
    doc.add_heading("Required Mathematician Deliverables (Game Description Section 16)", level=1)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "#"
    hdr[1].text = "Required Deliverable"
    hdr[2].text = "Provided Artifact(s)"
    hdr[3].text = "Status"

    rows = [
        (
            "1",
            "Final math report with model version",
            "Treasure_Dice_Phase_1_Math_Report.md",
            "Provided",
        ),
        (
            "2",
            "Stake Engine Math SDK package for eight modes",
            "math-sdk/games/treasure_dice",
            "Provided",
        ),
        (
            "3",
            "Approved regular routes paytable",
            "Treasure_Dice_Phase_1_Math_Report.md; treasure_dice_phase3_mapping_spec.json",
            "Provided",
        ),
        (
            "4",
            "Approved Bonus Buy paytable and Treasure Hunt description",
            "Treasure_Dice_Phase_2_Bonus_Report.md; treasure_dice_phase3_mapping_spec.json",
            "Provided",
        ),
        (
            "5",
            "index.json, lookUpTable.csv, compressed gameplay outcomes",
            "math-sdk/games/treasure_dice/library/publish_files",
            "Provided",
        ),
        (
            "6",
            "Max win multiplier for each mode",
            "treasure_dice_phase5_validation_summary.json; config_fe_treasure_dice.json",
            "Provided",
        ),
        (
            "7",
            "Simulation report",
            "Treasure_Dice_Phase_5_Simulation_Report.md",
            "Provided",
        ),
        (
            "8",
            "Volatility / variance report",
            "Treasure_Dice_Phase_5_Volatility_Report.md",
            "Provided",
        ),
        (
            "9",
            "Precision and rounding rules",
            "Treasure_Dice_Phase_1_Math_Report.md; Treasure_Dice_Phase_5_Replay_Audit_Report.md",
            "Provided (policy closure pending)",
        ),
        (
            "10",
            "Exposure table after receiving platform cap (if required)",
            "Treasure_Dice_Phase_2_Bonus_Report.md",
            "Conditional (awaiting cap input)",
        ),
        (
            "11",
            "Presentation event mapping for frontend and replay",
            "Treasure_Dice_Phase_3_Event_Mapping_Report.md; treasure_dice_phase3_mapping_spec.json",
            "Provided",
        ),
        (
            "12",
            "Brief risk assessment",
            "Treasure_Dice_Phase_6_Risk_Assessment.md",
            "Provided",
        ),
    ]

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value

    checks = validation_data.get("checks", {})
    all_checks_ok = all(
        checks.get(k, False)
        for k in [
            "all_weight_checks_pass",
            "all_lookup_checks_pass",
            "all_max_win_checks_pass",
            "all_replay_checks_pass",
        ]
    )
    doc.add_paragraph(
        f"Core package integrity checks: {'PASS' if all_checks_ok else 'FAIL'} "
        "(weight, lookup, max-win, replay)."
    )


def add_mode_table(doc: Document, validation_data: dict) -> None:
    doc.add_heading("Mode Summary and Max Win Multipliers", level=1)

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Mode"
    hdr[1].text = "Cost"
    hdr[2].text = "Theoretical RTP"
    hdr[3].text = "Simulated RTP"
    hdr[4].text = "Configured Max Win"
    hdr[5].text = "Replay Check"

    for mode in validation_data.get("modes", []):
        cells = table.add_row().cells
        cells[0].text = str(mode.get("mode_id", ""))
        cells[1].text = str(mode.get("cost", ""))
        cells[2].text = f"{float(mode.get('theoretical_rtp', 0.0)):.4f}"
        cells[3].text = f"{float(mode.get('simulated_rtp', 0.0)):.4f}"
        cells[4].text = str(mode.get("configured_max_win", ""))
        cells[5].text = "PASS" if mode.get("replay_check_pass") else "FAIL"


def add_publication_status(doc: Document) -> None:
    doc.add_heading("Publication Readiness", level=1)
    bullets = [
        "Status: Conditionally Ready.",
        "Blocking closure item: low-denomination regular-route rounding policy decision.",
        "Open dependency: operator payout-cap input if exposure enforcement is required.",
        "Gate 6 recommendation: Conditional Ready.",
    ]
    for line in bullets:
        doc.add_paragraph(line, style="List Bullet")


def main() -> None:
    validation_data = {}
    if VALIDATION_JSON.exists():
        validation_data = json.loads(VALIDATION_JSON.read_text(encoding="utf-8"))

    doc = Document()

    title = doc.add_heading("Treasure Dice - Mathematician Deliverables Package", level=0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    subtitle = doc.add_paragraph(f"Generated on {date.today().isoformat()}")
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_paragraph(
        "This document consolidates all required mathematician artifacts listed in "
        "Game Description Section 16, with references to produced reports, SDK package outputs, "
        "validation evidence, and publication-readiness status."
    )

    add_deliverables_table(doc, validation_data)
    add_mode_table(doc, validation_data)

    doc.add_heading("Phase and Governance Artifacts", level=1)
    phase_docs = [
        "Treasure_Dice_Mathematician_Handoff.md",
        "Treasure_Dice_Phase_1_Math_Report.md",
        "Treasure_Dice_Phase_2_Bonus_Report.md",
        "Treasure_Dice_Phase_3_Event_Mapping_Report.md",
        "Treasure_Dice_Phase_4_Package_Completeness_Report.md",
        "Treasure_Dice_Phase_5_Simulation_Report.md",
        "Treasure_Dice_Phase_5_Volatility_Report.md",
        "Treasure_Dice_Phase_5_Replay_Audit_Report.md",
        "Treasure_Dice_Phase_5_Final_Validation_Summary.md",
        "Treasure_Dice_Phase_6_Delivery_Manifest.md",
        "Treasure_Dice_Phase_6_Risk_Assessment.md",
        "Treasure_Dice_Phase_6_Acceptance_Checklist.md",
        "Treasure_Dice_Phase_6_Publication_Status_Note.md",
    ]
    for path in phase_docs:
        doc.add_paragraph(path, style="List Bullet")

    doc.add_heading("Machine-Readable Artifacts", level=1)
    machine_files = [
        "treasure_dice_phase3_mapping_spec.json",
        "treasure_dice_phase5_validation_summary.json",
        "treasure_dice_phase5_validate.py",
        "treasure_dice_phase1_simulator.py",
        "treasure_dice_phase2_simulator.py",
    ]
    for path in machine_files:
        doc.add_paragraph(path, style="List Bullet")

    add_file_list(doc, "Publish Files", PUBLISH_DIR)
    add_file_list(doc, "Lookup Table Files", LOOKUP_DIR)
    add_file_list(doc, "Config Files", CONFIG_DIR)
    add_file_list(doc, "Force Files", FORCE_DIR)

    add_publication_status(doc)

    doc.save(OUT_DOCX)
    print(f"Created: {OUT_DOCX}")


if __name__ == "__main__":
    main()
