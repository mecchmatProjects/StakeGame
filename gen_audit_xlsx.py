"""Generate AzteckPlinkoSet_audit.xlsx with calculated parameters."""
from __future__ import annotations

import json
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def create_audit_worksheet(
    wb: Workbook,
    mode_name: str,
    cost: float,
    lookup_file: Path,
) -> None:
    """Create a worksheet for a single mode with calculated parameters."""
    
    # Clean sheet name (max 31 chars for Excel)
    sheet_name = mode_name[:31]
    ws = wb.create_sheet(sheet_name)
    
    # Read lookup data (headerless CSV: sim_id,weight,payout_raw)
    outcomes = []
    with open(lookup_file, 'r', encoding='utf-8', newline='') as f:
        reader = csv.reader(f)
        for row in reader:
            if len(row) != 3:
                continue
            try:
                sim_id = int(row[0])
                weight = int(row[1])
                payout_raw = int(row[2])
            except ValueError:
                continue
            outcomes.append({
                'sim_id': sim_id,
                'weight': weight,
                'payout_raw': payout_raw,
            })
    
    # === Header Section ===
    ws['A1'] = f'{mode_name} Lookup Detail'
    ws['A1'].font = Font(bold=True, size=12)
    
    ws['A2'] = 'Derived from lookup rows'
    ws['A2'].font = Font(italic=True)
    
    # Metadata row
    ws['A3'] = 'Cost'
    ws['B3'] = cost
    ws['C3'] = 'Target RTP'
    ws['D3'] = 0.9605
    ws['E3'] = 'RTP Tol'
    ws['F3'] = 0.01
    ws['G3'] = 'Payout Scale'
    ws['H3'] = 100
    
    # Apply formatting
    for col in ['A', 'C', 'E', 'G']:
        ws[f'{col}3'].font = Font(bold=True)
    for col in ['B', 'D', 'F', 'H']:
        ws[f'{col}3'].number_format = '0.0000' if col in ['D', 'F'] else '0'
    
    # Empty row
    ws['A4'] = None
    
    # Rows header
    ws['A5'] = 'Rows'
    ws['A5'].font = Font(bold=True)
    
    # Column headers
    headers = ['sim_id', 'weight', 'payout_raw', 'payout_scaled', 'probability', 'weighted_payout', 'cumulative_probability']
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    
    # Data rows
    summary_row = len(outcomes) + 9
    total_weight_row = summary_row + 1

    for row_idx, outcome in enumerate(outcomes, 7):
        sim_id = outcome['sim_id']
        weight = outcome['weight']
        payout_raw = outcome['payout_raw']
        
        # sim_id
        ws.cell(row=row_idx, column=1).value = sim_id
        
        # weight
        ws.cell(row=row_idx, column=2).value = weight
        
        # payout_raw
        ws.cell(row=row_idx, column=3).value = payout_raw
        
        # payout_scaled (formula or value)
        payout_scaled_cell = ws.cell(row=row_idx, column=4)
        payout_scaled_cell.value = f'=C{row_idx}/$H$3'
        payout_scaled_cell.number_format = '0.00'
        
        # probability (formula with dynamic total-weight divisor)
        prob_cell = ws.cell(row=row_idx, column=5)
        prob_cell.value = f'=IF($B${total_weight_row}>0,B{row_idx}/$B${total_weight_row},0)'
        prob_cell.number_format = '0.0000'
        
        # weighted_payout (formula)
        weighted_cell = ws.cell(row=row_idx, column=6)
        weighted_cell.value = f'=D{row_idx}*E{row_idx}'
        weighted_cell.number_format = '0.0000'
        
        # cumulative_probability (formula)
        cum_cell = ws.cell(row=row_idx, column=7)
        if row_idx == 7:
            cum_cell.value = f'=E{row_idx}'
        else:
            cum_cell.value = f'=G{row_idx - 1}+E{row_idx}'
        cum_cell.number_format = '0.0000'
    
    # Summary section
    ws[f'A{summary_row}'] = 'Summary'
    ws[f'A{summary_row}'].font = Font(bold=True)
    
    ws[f'A{summary_row + 1}'] = 'Total Weight'
    ws[f'B{summary_row + 1}'] = f'=SUM(B7:B{len(outcomes) + 6})'
    
    ws[f'A{summary_row + 2}'] = 'Expected Value'
    ws[f'B{summary_row + 2}'] = f'=SUMPRODUCT(C7:C{len(outcomes) + 6},B7:B{len(outcomes) + 6})/B{summary_row + 1}'
    ws[f'B{summary_row + 2}'].number_format = '0.0000'
    
    ws[f'A{summary_row + 3}'] = 'RTP (vs Cost)'
    ws[f'B{summary_row + 3}'] = f'=B{summary_row + 2}/100/$B$3'
    ws[f'B{summary_row + 3}'].number_format = '0.0000%'

    ws[f'A{summary_row + 4}'] = 'RTP Check'
    ws[f'B{summary_row + 4}'] = f'=IF(AND(B{summary_row + 3}>=$D$3-$F$3, B{summary_row + 3}<=$D$3+$F$3), "PASS", "FAIL")'
    ws[f'B{summary_row + 4}'].font = Font(bold=True)
    
    # Column widths
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 18


def main() -> int:
    """Generate audit workbook for Azteck Plinko Set."""
    
    # Read index from Azteck_plinko_set
    index_path = Path('Azteck_plinko_set/artifacts/publish_files/index.json')
    with open(index_path, 'r', encoding='utf-8') as f:
        index = json.load(f)
    
    # Create workbook
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet
    
    # Generate worksheet for each mode
    modes_processed = 0
    for mode in index['modes']:
        mode_name = mode['name']
        cost = mode['cost']
        lookup_file = Path('Azteck_plinko_set/artifacts/publish_files') / mode['weights']
        
        try:
            create_audit_worksheet(wb, mode_name, cost, lookup_file)
            modes_processed += 1
            print(f'✓ {mode_name}')
        except Exception as e:
            print(f'✗ {mode_name}: {e}')
    
    # Save workbook
    output_path = Path('AzteckPlinkoSet_audit.xlsx')
    wb.save(output_path)

    # Save alias with requested spelling
    output_path_alias = Path('AztecPlinkoSet_audit.xlsx')
    wb.save(output_path_alias)
    
    print(f'\nCreated {output_path}')
    print(f'Modes: {modes_processed}/{len(index["modes"])}')
    print(f'File size: {output_path.stat().st_size / 1024:.2f} KB')
    print(f'Alias created: {output_path_alias}')
    
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
